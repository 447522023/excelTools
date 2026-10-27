"""把对比结果导出为 xlsx(带状态/差异单元格着色)或 csv(utf-8-sig,Excel 可直接打开中文)"""
import csv
import io
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .comparer import (
    STATUS_CHANGED,
    STATUS_ONLY_A,
    STATUS_ONLY_B,
    STATUS_SAME,
    CompareOutcome,
)

_STATUS_FILL = {
    STATUS_ONLY_A: PatternFill("solid", fgColor="FFF3E0"),
    STATUS_ONLY_B: PatternFill("solid", fgColor="E3F2FD"),
    STATUS_CHANGED: PatternFill("solid", fgColor="FFEBEE"),
    STATUS_SAME: PatternFill("solid", fgColor="E8F5E9"),
}
_DIFF_FILL = PatternFill("solid", fgColor="FFEBEE")
_HEADER_FILL = PatternFill("solid", fgColor="F0F2F5")


def to_xlsx(outcome: CompareOutcome) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "对比结果"
    columns = outcome.columns

    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 差异列成对出现:"xxx（A）" 与 "xxx（B）",导出时逐对比较,不一致的单元格标红
    pair_right: dict[int, int] = {}
    col_index = {name: i for i, name in enumerate(columns)}
    for i, name in enumerate(columns):
        if name.endswith("（A）") and f"{name[:-3]}（B）" in col_index:
            pair_right[i] = col_index[f"{name[:-3]}（B）"]

    for row in outcome.rows:
        ws.append([row.get(c, "") for c in columns])
        r = ws.max_row
        status = row.get("状态", "")
        status_cell = ws.cell(row=r, column=1)
        if status in _STATUS_FILL:
            status_cell.fill = _STATUS_FILL[status]
            status_cell.alignment = Alignment(horizontal="center")
        if status == STATUS_CHANGED:
            for left_col, right_col in pair_right.items():
                lv = ws.cell(row=r, column=left_col + 1).value
                rv = ws.cell(row=r, column=right_col + 1).value
                if lv != rv:
                    ws.cell(row=r, column=left_col + 1).fill = _DIFF_FILL
                    ws.cell(row=r, column=right_col + 1).fill = _DIFF_FILL

    for i, name in enumerate(columns, start=1):
        width = max(len(str(name)), *(len(str(row.get(name, ""))) for row in outcome.rows[:200])) + 4
        ws.column_dimensions[get_column_letter(i)].width = min(width, 45)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(outcome: CompareOutcome) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=outcome.columns, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(outcome.rows)
    return buf.getvalue().encode("utf-8-sig")


def content_disposition(filename: str) -> str:
    return f"attachment; filename*=UTF-8''{quote(filename)}"
