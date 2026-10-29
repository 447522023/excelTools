"""统一读取 xls / xlsx / xlsm / csv 为"全部规范化字符串"的 DataFrame。

规范化规则:空值 -> "";整数浮点 1.0 -> "1";其余 str() 后去首尾空白。
这样 1 与 "1"、1.0 与 1 在对比时视为相同。

合并表头支持:所选表头行落在合并区域内时,自动识别整条"表头带"(可跨多行):
- 纵向合并(如"编号"占 2 行):取合并单元格的值作为列名;
- 横向合并(如"基本信息"盖住 2 列):值铺到覆盖的每一列,再与下一级的子表头拼接;
- 多行表头按"上级-下级"拼接成唯一列名(如"基本信息-姓名");
- 数据一律从表头带最后一行之下开始,两表表头行数不一致也能各自正确解析。

文件类型按内容嗅探决定(PK 头 -> xlsx 家族,OLE2 头 -> xls,其余 -> csv),
扩展名只是初筛,改名/误存的文件也能正确读取。
"""
import math
import re
from functools import lru_cache
from pathlib import Path

import pandas as pd
from charset_normalizer import from_path
from openpyxl import load_workbook

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".csv"}

# 纯数字且小数部分全为 0 的字符串,如 "10.0",去掉浮点尾巴变 "10"(保留前导零,如 "00123.0" -> "00123")
_TRAILING_ZERO_FLOAT = re.compile(r"^-?\d+\.0+$")

# pandas 对空表头单元格生成的占位名,如 "Unnamed: 2"
_UNNAMED = re.compile(r"^Unnamed: \d+$")

# 表头带最多覆盖的行数,防止把数据区的大合并误认成表头
_MAX_HEADER_BAND = 10

# 合并区域统一为 0 起半开区间 [r1, r2) x [c1, c2)
_Merge = tuple[int, int, int, int]

# OOXML 命名空间(workbook.xml 的 sheet 节点与关系 id)
_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_RNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def normalize_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        return str(int(v)) if v.is_integer() else str(v)
    s = str(v).strip()
    if _TRAILING_ZERO_FLOAT.fullmatch(s):
        return s.split(".")[0]
    return s


def _dedupe_columns(cols) -> list[str]:
    seen: dict[str, int] = {}
    out: list[str] = []
    for i, c in enumerate(cols):
        name = str(c).strip()
        if not name or _UNNAMED.fullmatch(name):
            # 空表头按列位置命名,两表都有空表头时也能按位置对上
            name = f"未命名{i + 1}"
        if name in seen:
            seen[name] += 1
            out.append(f"{name}.{seen[name]}")
        else:
            seen[name] = 1
            out.append(name)
    return out


def _normalize_df(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = _dedupe_columns(df.columns)
    elementwise = df.map if hasattr(df, "map") else df.applymap
    return elementwise(normalize_value)


def _detect_encoding(path: Path) -> str | None:
    """探测 CSV 编码。

    utf-8 / BOM 先严格验证;其次取探测结果中可信度大于 0 的那个。
    全部不可信时(常见于只有几个汉字的小文件:GBK 与韩文 cp949 字节高度重叠,
    探测器只能给出 0 置信度的猜测),面向中文场景优先严格试 GBK,
    再不行才退回探测器的最佳猜测。
    """
    raw = path.read_bytes()[:65536]
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    try:
        raw.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        matches = list(from_path(str(path)))
    except Exception:
        matches = []
    confident = next(
        (m.encoding for m in matches if getattr(m, "coherence", 0) > 0), None
    )
    if confident:
        return confident
    for enc in ("gbk", *(m.encoding for m in matches)):
        try:
            raw.decode(enc)
            return enc
        except (UnicodeDecodeError, LookupError):
            continue
    return None  # 调用方以 utf-8 兜底


def _cell(grid: list[list], r: int, c: int):
    if 0 <= r < len(grid) and 0 <= c < len(grid[r]):
        return grid[r][c]
    return None


def _xlsx_merges(path: Path) -> dict[str, list[_Merge]]:
    """从 zip 内 XML 直接提取合并区域(比整本 load_workbook 轻量得多,大文件友好)"""
    import xml.etree.ElementTree as ET
    import zipfile

    from openpyxl.utils import range_boundaries

    try:
        with zipfile.ZipFile(path) as z:
            wb = ET.fromstring(z.read("xl/workbook.xml"))
            rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            rid_to_target = {r.get("Id"): r.get("Target") for r in rels}
            sheets_el = wb.find(f"{_NS}sheets")
            out: dict[str, list[_Merge]] = {}
            for sh in sheets_el.iter(f"{_NS}sheet") if sheets_el is not None else ():
                rid = sh.get(f"{_RNS}id", "")
                target = (rid_to_target.get(rid) or "").lstrip("/")
                if not target:
                    continue
                xml_path = target if target.startswith("xl/") else f"xl/{target}"
                if xml_path not in z.namelist():
                    continue
                data = z.read(xml_path)
                merges: list[_Merge] = []
                for m in re.finditer(rb'<mergeCell ref="([A-Z0-9]+:[A-Z0-9]+)"', data):
                    c1, r1, c2, r2 = range_boundaries(m.group(1).decode())
                    merges.append((r1 - 1, r2, c1 - 1, c2))
                out[sh.get("name", "")] = merges
            return out
    except Exception:
        return {}


def _xls_merges(path: Path) -> dict[str, list[_Merge]]:
    try:
        import xlrd

        book = xlrd.open_workbook(str(path), formatting_info=True)
    except Exception:
        # 个别老格式读不到格式信息时,按"无合并单元格"处理
        return {}
    # xlrd 的 merged_cells 本身就是 0 起半开区间
    return {sh.name: [tuple(m) for m in sh.merged_cells] for sh in book.sheets()}


def _merge_map(path: Path) -> dict[str, list[_Merge]]:
    kind = sniff_kind(path)
    if kind == "xls":
        return _xls_merges(path)
    if kind == "xlsx":
        return _xlsx_merges(path)
    return {}


def _grid_rows(df: pd.DataFrame) -> list[list]:
    return [list(row) for row in df.itertuples(index=False, name=None)]


def _header_band(grid: list[list], merges: list[_Merge], hr: int) -> tuple[int, int]:
    """表头带 [top, bottom](0 起闭区间):从所选行 hr 出发,沿"跨过带边界且有值"的
    合并单元格向上/向下扩展。数据区/标题区的合并不会跨界,天然不受影响。
    """
    top, bottom = hr, hr
    while True:
        ntop, nbottom = top, bottom
        for r1, r2, c1, _c2 in merges:
            if not normalize_value(_cell(grid, r1, c1)):
                continue  # 空的合并区域(如纯格式合并)不参与表头带识别
            if r1 < top < r2 and bottom - r1 + 1 <= _MAX_HEADER_BAND:
                ntop = min(ntop, r1)  # 合并从带上方跨进来 -> 表头带向上扩
            if r1 <= bottom < r2 and r2 - top <= _MAX_HEADER_BAND:
                nbottom = max(nbottom, r2 - 1)  # 合并从带内跨到下方 -> 表头带向下扩
        if (ntop, nbottom) == (top, bottom):
            return top, bottom
        top, bottom = ntop, nbottom


def _anchor_values(
    grid: list[list], merges: list[_Merge], top: int, bottom: int, ncols: int
) -> dict[tuple[int, int], str]:
    """表头带内每个格子的"合并锚点值":格子属于某合并区域时取该区域左上角的值"""
    anchor: dict[tuple[int, int], str] = {}
    for r1, r2, c1, c2 in merges:
        value = normalize_value(_cell(grid, r1, c1))
        if not value:
            continue
        for r in range(max(r1, top), min(r2, bottom + 1)):
            for c in range(max(c1, 0), min(c2, ncols)):
                anchor[(r, c)] = value
    return anchor


def _frame_from_grid(grid: list[list], merges: list[_Merge], header_row: int) -> pd.DataFrame:
    """原始网格 + 合并信息 + 表头行(1 起)-> DataFrame(列名为拼接后的表头)"""
    if not grid:  # 空工作表:不报错,交由上层决定(坐标对比可用,表头解析会给明确提示)
        return _normalize_df(pd.DataFrame())
    if header_row < 1 or header_row > len(grid):
        raise ValueError(f"表头行 {header_row} 超出表格范围(共 {len(grid)} 行)")
    hr = header_row - 1
    ncols = max((len(r) for r in grid), default=0)
    top, bottom = _header_band(grid, merges, hr)
    anchor = _anchor_values(grid, merges, top, bottom, ncols)

    names: list[str] = []
    fallback_cols: set[int] = set()  # 表头完全为空的列,若数据也全空则整列丢弃
    for c in range(ncols):
        parts: list[str] = []
        for r in range(top, bottom + 1):
            v = anchor.get((r, c)) or normalize_value(_cell(grid, r, c))
            if v and v not in parts:
                parts.append(v)
        if parts:
            names.append("-".join(parts))
        else:
            names.append(f"未命名{c + 1}")
            fallback_cols.add(c)

    # 数据从表头带下方开始;剔除整行全空的行(合并表头下方常有多余空行)
    data = [
        row for row in grid[bottom + 1 :] if any(normalize_value(v) for v in row)
    ]
    # 先规范化再判断幻影列:None/NaN 都归一为 "",全空列才能被正确剔除
    df = _normalize_df(pd.DataFrame(data, columns=names))
    keep = [
        name
        for i, name in enumerate(names)
        if i not in fallback_cols or not df[name].eq("").all()
    ]
    return df[keep]


def sniff_kind(path: Path) -> str:
    """按文件头嗅探真实类型。扩展名只作参考,误存/改名的文件也能读对:
    PK zip 头 -> xlsx 家族(xlsx/xlsm);OLE2 头 -> xls;其余 -> csv 文本。
    """
    try:
        with path.open("rb") as f:
            head = f.read(8)
    except OSError:
        return "csv"
    if head[:4] in (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"):
        return "xlsx"
    if head == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    return "csv"


def _xlsx_grids(path: Path) -> dict[str, pd.DataFrame]:
    """openpyxl 直读网格(值层)。不经 pandas 的列推断:
    某列全是"无缓存值的公式"时 pandas 会把整列丢掉,坐标对比不允许丢列。
    """
    wb = load_workbook(path.open("rb"), read_only=True, data_only=True)  # 传文件对象:不按扩展名误判
    try:
        out: dict[str, pd.DataFrame] = {}
        for ws in wb.worksheets:
            rows = [list(row) for row in ws.iter_rows(values_only=True)]
            out[ws.title] = pd.DataFrame(rows)
        return out
    finally:
        wb.close()


@lru_cache(maxsize=4)
def _raw_sheets_cached(path_str: str, mtime: float) -> dict[str, pd.DataFrame]:
    """原始网格缓存(含表头行,未做任何表头解析)。

    上传(列信息/预览)与对比(表头解析)共用同一次磁盘读取,
    大文件不必重复解析;文件在存储层不可变,以 mtime 兜底失效。
    """
    path = Path(path_str)
    kind = sniff_kind(path)
    if kind == "csv":
        encoding = _detect_encoding(path) or "utf-8"
        try:
            return {
                "CSV": pd.read_csv(
                    path, encoding=encoding, sep=None, engine="python",
                    dtype=object, keep_default_na=False, header=None,
                )
            }
        except pd.errors.EmptyDataError:  # 完全空白的 csv 不报错,按空表处理
            return {"CSV": pd.DataFrame()}
    if kind == "xlsx":
        return _xlsx_grids(path)
    return pd.read_excel(path, sheet_name=None, engine="xlrd", dtype=object, header=None)


def _raw_sheets(path: Path) -> dict[str, pd.DataFrame]:
    stat = path.stat()
    return _raw_sheets_cached(str(path), stat.st_mtime)


@lru_cache(maxsize=2)
def _frames_cached(
    path_str: str, mtime: float, sheet: str | None, header_row: int
) -> dict[str, pd.DataFrame]:
    """表头解析结果缓存:上传后紧接着的对比/导出不再重复解析同一文件。

    返回值视为只读(调用方不得原地修改 DataFrame)。
    """
    path = Path(path_str)
    grids = _raw_sheets(path)
    merges = _merge_map(path)
    if sheet is None:
        names = list(grids)
    else:
        if sheet not in grids:
            raise ValueError(f"没有名为“{sheet}”的工作表")
        names = [sheet]
    return {
        name: _frame_from_grid(_grid_rows(grids[name]), merges.get(name, []), header_row)
        for name in names
    }


def _read_excel_sheets(
    path: Path, sheet: str | None, header_row: int
) -> dict[str, pd.DataFrame]:
    stat = path.stat()
    return _frames_cached(str(path), stat.st_mtime, sheet, max(1, header_row))


def _read_csv(path: Path, header_row: int) -> pd.DataFrame:
    # 统一走网格解析管线,与 excel 的列名/空行处理完全一致(跨格式对比行为一致)
    return _frame_from_grid(_grid_rows(_raw_sheets(path)["CSV"]), [], max(1, header_row))


def _check_supported(path: Path) -> None:
    suffix = Path(path).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"不支持的文件类型:{suffix},仅支持 {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )


def read_tables(path: Path, header_row: int = 1) -> dict[str, pd.DataFrame]:
    """读取整个文件,返回 {sheet名: DataFrame}。csv 视为单个名为 CSV 的表。

    header_row:表头所在行数(1 起);该行所在的合并表头带之上的内容会被忽略。
    """
    _check_supported(path)
    return _read_excel_sheets(path, None, max(1, header_row))


def read_table(path: Path, sheet: str, header_row: int = 1) -> pd.DataFrame:
    """只读取指定工作表(表头行可指定),避免其他工作表的解析问题影响本次请求"""
    _check_supported(path)
    if sniff_kind(path) == "csv":
        if sheet != "CSV":
            raise ValueError(f"CSV 文件只有一个名为 CSV 的表,没有“{sheet}”")
        return _read_csv(path, max(1, header_row))
    return _read_excel_sheets(path, sheet, max(1, header_row))[sheet]


def sheet_names(path: Path) -> list[str]:
    """只取工作表名(不解析数据),用于提示两表各有哪些 sheet"""
    kind = sniff_kind(path)
    if kind == "csv":
        return ["CSV"]
    if kind == "xls":
        import xlrd

        return xlrd.open_workbook(str(path), on_demand=True).sheet_names()
    import xml.etree.ElementTree as ET
    import zipfile

    with zipfile.ZipFile(path) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        sheets_el = wb.find(f"{_NS}sheets")
        return [
            s.get("name", "") for s in (sheets_el.iter(f"{_NS}sheet") if sheets_el is not None else ())
        ]


def read_grid(path: Path, sheet: str | None = None) -> dict[str, list[list[str]]]:
    """规范化文本网格(不做任何表头假设),坐标对比模式用。尾部全空行忽略。"""
    grids = _raw_sheets(path)
    out: dict[str, list[list[str]]] = {}
    for name, df in grids.items():
        if sheet is not None and name != sheet:
            continue
        grid = [
            [normalize_value(v) for v in row]
            for row in df.itertuples(index=False, name=None)
        ]
        while grid and all(v == "" for v in grid[-1]):
            grid.pop()
        out[name] = grid
    if sheet is not None and sheet not in out:
        raise ValueError(f"没有名为“{sheet}”的工作表")
    return out


def read_formulas(path: Path) -> dict[str, dict[tuple[int, int], str]]:
    """公式文本 {sheet名: {(行0起, 列0起): 公式串}}。仅 xlsx 家族;xls/csv 无公式层。"""
    if sniff_kind(path) != "xlsx":
        return {}
    out: dict[str, dict[tuple[int, int], str]] = {}
    wb = load_workbook(path.open("rb"), read_only=True, data_only=False)  # 文件对象:防扩展名误判
    try:
        for ws in wb.worksheets:
            formulas: dict[tuple[int, int], str] = {}
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("="):
                        formulas[(cell.row - 1, cell.column - 1)] = v
            if formulas:
                out[ws.title] = formulas
    finally:
        wb.close()
    return out


def describe_file(
    path: Path, header_row: int = 1, max_rows: int = 8, max_cols: int = 20
) -> dict[str, dict]:
    """一次磁盘读取,产出每个 sheet 的行数/列名/原始预览(上传接口专用)"""
    _check_supported(path)
    grids = _raw_sheets(path)
    merges = _merge_map(path)
    out: dict[str, dict] = {}
    for name, df in grids.items():
        rows = _grid_rows(df)
        frame = _frame_from_grid(rows, merges.get(name, []), header_row)
        out[name] = {
            "row_count": len(frame),
            "columns": list(frame.columns),
            "raw_preview": [
                [normalize_value(v) for v in row[:max_cols]] for row in rows[:max_rows]
            ],
        }
    return out


def raw_preview(path: Path, max_rows: int = 8, max_cols: int = 20) -> dict[str, list[list[str]]]:
    """每个 sheet 前 N 行的原始文本(不做任何表头假设),供前端展示、帮助用户挑表头行"""
    _check_supported(path)
    out: dict[str, list[list[str]]] = {}
    for name, df in _raw_sheets(path).items():
        grid = []
        for row in df.itertuples(index=False, name=None):
            grid.append([normalize_value(v) for v in list(row)[:max_cols]])
        out[name] = grid[:max_rows]
    return out
