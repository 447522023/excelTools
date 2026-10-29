import base64
from pathlib import Path

import pytest

from app.services import ops
from app.storage import FileStore

DATA_DIR = Path(__file__).resolve().parent.parent.parent / "test_files"


@pytest.fixture()
def store(tmp_path):
    return FileStore(root=tmp_path / "uploads")


def _read(p: str) -> bytes:
    return (DATA_DIR / p).read_bytes()


def test_full_flow_upload_parse_compare_export(store):
    up_a = ops.upload_bytes(store, "a.xlsx", _read("员工名单A_10月.xlsx"))
    up_b = ops.upload_bytes(store, "b.csv", _read("员工名单B_10月.csv"))  # GBK
    assert up_b["sheets"][0]["columns"][0] == "编号"

    parsed = ops.parse_sheet(store, up_b["file_id"], "CSV", header_row=1)
    assert parsed["row_count"] == 11

    payload = {
        "left_file_id": up_a["file_id"], "left_sheet": "10月名单", "left_header_row": 1,
        "right_file_id": up_b["file_id"], "right_sheet": "CSV", "right_header_row": 1,
        "left_keys": ["编号"], "right_keys": ["编号"], "mode": "diff",
        "compare_columns": None, "page": 1, "page_size": 0,
    }
    res = ops.compare_page(store, payload)
    assert res["total"] == 7 and len(res["rows"]) == 7
    assert res["mode"] == "diff"

    filename, data = ops.export_result(store, {**payload, "format": "xlsx"})
    assert filename.endswith(".xlsx") and len(data) > 1000


def test_upload_rejects_bad_type(store):
    with pytest.raises(ops.OpError) as e:
        ops.upload_bytes(store, "x.txt", b"hello")
    assert "不支持的文件类型" in e.value.message


def test_compare_key_mismatch(store):
    up_a = ops.upload_bytes(store, "a.xlsx", _read("员工名单A_10月.xlsx"))
    up_b = ops.upload_bytes(store, "b.csv", _read("员工名单B_10月.csv"))
    payload = {
        "left_file_id": up_a["file_id"], "left_sheet": "10月名单", "left_header_row": 1,
        "right_file_id": up_b["file_id"], "right_sheet": "CSV", "right_header_row": 1,
        "left_keys": ["编号", "姓名"], "right_keys": ["编号"], "mode": "diff",
        "compare_columns": None, "page": 1, "page_size": 0,
    }
    with pytest.raises(ops.OpError) as e:
        ops.compare_page(store, payload)
    assert "数量不一致" in e.value.message


def test_base64_roundtrip_matches_bridge_usage(store):
    """离线客户端桥接层用 base64 传文件,验证该路径等价"""
    raw = _read("员工名单D_带标题行.xlsx")
    up = ops.upload_bytes(store, "d.xlsx", base64.b64decode(base64.b64encode(raw)))
    parsed = ops.parse_sheet(store, up["file_id"], "名单", header_row=3)
    assert parsed["columns"] == ["编号", "姓名", "部门", "基本工资", "入职日期"]
    assert parsed["row_count"] == 6


def test_cell_mode_end_to_end(store, tmp_path):
    """坐标模式全链路:无需关键列,输出差异行(含公式层),可导出"""
    from openpyxl import Workbook

    def make(path, b1_formula):
        wb = Workbook()
        ws = wb.active
        ws.title = "S"
        ws["A1"], ws["A2"] = 10, 20
        ws["B1"], ws["B2"] = b1_formula, "=A2+1"
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)

    old, new = tmp_path / "old.xlsx", tmp_path / "new.xlsx"
    make(old, "=A1+1")
    make(new, "=1+A1")  # 公式文本变了,结果不变
    up_a = ops.upload_bytes(store, old.name, old.read_bytes())
    up_b = ops.upload_bytes(store, new.name, new.read_bytes())
    payload = {
        "left_file_id": up_a["file_id"], "left_sheet": "S", "left_header_row": 1,
        "right_file_id": up_b["file_id"], "right_sheet": "S", "right_header_row": 1,
        "left_keys": [], "right_keys": [], "mode": "cell",
        "compare_columns": None, "page": 1, "page_size": 0,
    }
    res = ops.compare_page(store, payload)
    assert res["mode"] == "cell"
    assert res["total"] == 1  # 必须真的输出差异行(回归:曾因模式传错输出 0 行)
    assert res["summary"]["公式差异格"] == 1
    assert "[公式]" in res["rows"][0]["差异明细"]

    filename, data = ops.export_result(store, {**payload, "format": "xlsx"})
    assert "坐标对比" in filename and len(data) > 1000


def test_cell_mode_empty_sheet_side(store, tmp_path):
    """一侧工作表完全为空:上传不报错,坐标对比输出全部行仅表A"""
    from openpyxl import Workbook

    old = tmp_path / "old.xlsx"
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = "编号"
    wb.active["A2"] = "K1"
    wb.save(old)
    new = tmp_path / "new.xlsx"
    wb2 = Workbook()
    wb2.active.title = "S"
    wb2.save(new)

    up_a = ops.upload_bytes(store, "old.xlsx", old.read_bytes())
    up_b = ops.upload_bytes(store, "new.xlsx", new.read_bytes())
    assert up_b["sheets"][0]["row_count"] == 0  # 空表可上传

    res = ops.compare_page(store, {
        "left_file_id": up_a["file_id"], "left_sheet": "S", "left_header_row": 1,
        "right_file_id": up_b["file_id"], "right_sheet": "S", "right_header_row": 1,
        "left_keys": [], "right_keys": [], "mode": "cell",
        "compare_columns": None, "page": 1, "page_size": 0,
    })
    assert res["summary"]["仅表A"] == 2  # 表头行 + 1 行数据


def test_key_mode_still_requires_keys(store):
    up_a = ops.upload_bytes(store, "a.xlsx", _read("员工名单A_10月.xlsx"))
    up_b = ops.upload_bytes(store, "b.csv", _read("员工名单B_10月.csv"))
    with pytest.raises(ops.OpError) as e:
        ops.compare_page(store, {
            "left_file_id": up_a["file_id"], "left_sheet": "10月名单", "left_header_row": 1,
            "right_file_id": up_b["file_id"], "right_sheet": "CSV", "right_header_row": 1,
            "left_keys": [], "right_keys": [], "mode": "diff",
            "compare_columns": None, "page": 1, "page_size": 0,
        })
    assert "关键列" in e.value.message


def test_cell_mode_response_passes_api_schema(store, tmp_path):
    """回归:坐标模式结果必须能通过 HTTP 层的 CompareResponse 校验(全字符串行)"""
    from openpyxl import Workbook

    old, new = tmp_path / "old.xlsx", tmp_path / "new.xlsx"
    for path, name in ((old, "甲"), (new, "乙")):
        wb = Workbook()
        wb.active.title = "S"
        wb.active["A1"] = "值"
        wb.active["A2"] = name
        wb.save(path)
    up_a = ops.upload_bytes(store, old.name, old.read_bytes())
    up_b = ops.upload_bytes(store, new.name, new.read_bytes())
    res = ops.compare_page(store, {
        "left_file_id": up_a["file_id"], "left_sheet": "S", "left_header_row": 1,
        "right_file_id": up_b["file_id"], "right_sheet": "S", "right_header_row": 1,
        "left_keys": [], "right_keys": [], "mode": "cell",
        "compare_columns": None, "page": 1, "page_size": 0,
    })
    from app.schemas import CompareResponse

    CompareResponse(**res)  # 不抛异常即通过(与路由层同一模型)
